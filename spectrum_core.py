"""CSV loading, spectrum statistics, and export helpers.

The loader is intentionally tolerant of Vernier-style exports: it accepts
comma, tab, and semicolon delimiters, metadata before the table, repeated
wavelength columns, and Korean/English headers.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import csv
import math
import re
from typing import Iterable

import numpy as np


WAVELENGTH_WORDS = ("wavelength", "wave length", "lambda", "파장", "nm")


@dataclass(frozen=True)
class SpectrumSeries:
    label: str
    wavelength_nm: np.ndarray
    values: np.ndarray
    value_unit: str = ""


@dataclass(frozen=True)
class SpectrumFile:
    source: Path
    encoding: str
    delimiter: str
    series: tuple[SpectrumSeries, ...]


@dataclass(frozen=True)
class RegionSummary:
    center_nm: float
    half_width_nm: float
    count: int
    nearest_wavelength_nm: float
    nearest_value: float
    maximum_wavelength_nm: float
    maximum_value: float
    minimum_wavelength_nm: float
    minimum_value: float
    mean_value: float


def _decode_file(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1"), "latin-1"


def _choose_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines() if line.strip()][:50]
    sample = "\n".join(lines)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
    except csv.Error:
        counts = {delimiter: sample.count(delimiter) for delimiter in (",", "\t", ";")}
        return max(counts, key=counts.get)


def _number(value: str) -> float | None:
    text = value.strip().replace("\u00a0", "").replace(" ", "")
    if not text:
        return None
    if re.fullmatch(r"[+-]?(?:\d+,\d+)(?:[eE][+-]?\d+)?", text):
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _numeric_count(row: list[str]) -> int:
    return sum(_number(cell) is not None for cell in row)


def _find_data_start(rows: list[list[str]]) -> int:
    for index, row in enumerate(rows):
        if _numeric_count(row) < 2:
            continue
        following = rows[index : index + 3]
        if len(following) >= 2 and all(_numeric_count(item) >= 2 for item in following):
            return index
    raise ValueError("두 개 이상의 숫자 열이 연속된 스펙트럼 표를 찾지 못했습니다.")


def _header_for_column(rows: list[list[str]], data_start: int, column: int) -> str:
    candidates: list[str] = []
    for row in rows[max(0, data_start - 4) : data_start]:
        if column >= len(row):
            continue
        value = row[column].strip()
        if value and _number(value) is None:
            candidates.append(value)
    return " ".join(candidates[-2:]).strip() or f"열 {column + 1}"


def _is_wavelength_header(header: str) -> bool:
    lower = header.casefold()
    return any(word in lower for word in WAVELENGTH_WORDS)


def _looks_like_wavelength(values: np.ndarray, *, short_range_allowed: bool = False) -> bool:
    finite = values[np.isfinite(values)]
    minimum_count = 3 if short_range_allowed else 8
    minimum_span = 1 if short_range_allowed else 20
    if finite.size < minimum_count:
        return False
    diffs = np.diff(finite)
    monotonic_ratio = max(np.mean(diffs >= 0), np.mean(diffs <= 0)) if diffs.size else 0
    return (
        monotonic_ratio >= 0.9
        and float(np.nanmin(finite)) >= 150
        and float(np.nanmax(finite)) <= 2500
        and float(np.nanmax(finite) - np.nanmin(finite)) >= minimum_span
    )


def _value_unit(header: str) -> str:
    lower = header.casefold()
    if "absorb" in lower or "흡광" in lower:
        return "A.U."
    if "transmit" in lower or "투과" in lower or "%t" in lower:
        return "%T"
    if "intensity" in lower or "intens" in lower or "세기" in lower:
        return ""
    match = re.search(r"(?:\(|\[)([^)\]]+)(?:\)|\])", header)
    return match.group(1).strip() if match else ""


def _unique_label(label: str, existing: set[str]) -> str:
    base = label or "측정값"
    candidate = base
    number = 2
    while candidate in existing:
        candidate = f"{base} ({number})"
        number += 1
    return candidate


def load_spectrum_csv(filename: str | Path) -> SpectrumFile:
    path = Path(filename)
    if not path.is_file():
        raise FileNotFoundError(path)
    text, encoding = _decode_file(path)
    delimiter = _choose_delimiter(text)
    rows = [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError("CSV 파일이 비어 있습니다.")

    data_start = _find_data_start(rows)
    width = max(len(row) for row in rows[data_start:])
    headers = [_header_for_column(rows, data_start, column) for column in range(width)]

    columns: list[np.ndarray] = []
    for column in range(width):
        values = [
            _number(row[column]) if column < len(row) else None
            for row in rows[data_start:]
        ]
        columns.append(np.array([np.nan if value is None else value for value in values], dtype=float))

    wavelength_columns = [
        index
        for index, values in enumerate(columns)
        if _is_wavelength_header(headers[index])
        and _looks_like_wavelength(values, short_range_allowed=True)
    ]
    if not wavelength_columns:
        wavelength_columns = [
            index for index, values in enumerate(columns) if _looks_like_wavelength(values)
        ]
    if not wavelength_columns:
        raise ValueError("파장 열을 찾지 못했습니다. 파장 단위가 nm인지 확인하십시오.")

    series: list[SpectrumSeries] = []
    used_labels: set[str] = set()
    wavelength_set = set(wavelength_columns)
    for value_index, values in enumerate(columns):
        if value_index in wavelength_set:
            continue
        finite_count = int(np.isfinite(values).sum())
        if finite_count < 3:
            continue
        preceding = [index for index in wavelength_columns if index < value_index]
        wavelength_index = preceding[-1] if preceding else wavelength_columns[0]
        wavelengths = columns[wavelength_index]
        valid = np.isfinite(wavelengths) & np.isfinite(values)
        if int(valid.sum()) < 3:
            continue
        x = wavelengths[valid]
        y = values[valid]
        order = np.argsort(x, kind="stable")
        x = x[order]
        y = y[order]
        keep = np.concatenate(([True], np.diff(x) != 0))
        x = x[keep]
        y = y[keep]
        label = _unique_label(headers[value_index], used_labels)
        used_labels.add(label)
        series.append(SpectrumSeries(label, x, y, _value_unit(label)))

    if not series:
        raise ValueError("파장과 짝을 이루는 측정값 열을 찾지 못했습니다.")
    return SpectrumFile(path, encoding, delimiter, tuple(series))


def load_spectrum_xlsx(filename: str | Path) -> SpectrumFile:
    """Load a spectrum from the first worksheet containing a numeric table."""

    path = Path(filename)
    if not path.is_file():
        raise FileNotFoundError(path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX를 읽으려면 openpyxl이 필요합니다: python -m pip install openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    failures: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            try:
                dataset = _spectrum_file_from_rows(path, rows, "xlsx", f"xlsx:{sheet.title}")
                return dataset
            except ValueError as exc:
                failures.append(f"{sheet.title}: {exc}")
    finally:
        workbook.close()
    raise ValueError("XLSX에서 스펙트럼 표를 찾지 못했습니다. " + "; ".join(failures[:3]))


def _spectrum_file_from_rows(
    path: Path, rows: list[list[str]], encoding: str, delimiter: str
) -> SpectrumFile:
    if not rows:
        raise ValueError("파일이 비어 있습니다.")
    data_start = _find_data_start(rows)
    width = max(len(row) for row in rows[data_start:])
    headers = [_header_for_column(rows, data_start, column) for column in range(width)]
    columns: list[np.ndarray] = []
    for column in range(width):
        values = [_number(row[column]) if column < len(row) else None for row in rows[data_start:]]
        columns.append(np.array([np.nan if value is None else value for value in values], dtype=float))
    wavelength_columns = [
        index
        for index, values in enumerate(columns)
        if _is_wavelength_header(headers[index]) and _looks_like_wavelength(values, short_range_allowed=True)
    ]
    if not wavelength_columns:
        wavelength_columns = [index for index, values in enumerate(columns) if _looks_like_wavelength(values)]
    if not wavelength_columns:
        raise ValueError("파장 열을 찾지 못했습니다. 파장 단위가 nm인지 확인하십시오.")
    series: list[SpectrumSeries] = []
    used_labels: set[str] = set()
    wavelength_set = set(wavelength_columns)
    for value_index, values in enumerate(columns):
        if value_index in wavelength_set or int(np.isfinite(values).sum()) < 3:
            continue
        preceding = [index for index in wavelength_columns if index < value_index]
        wavelength_index = preceding[-1] if preceding else wavelength_columns[0]
        wavelengths = columns[wavelength_index]
        valid = np.isfinite(wavelengths) & np.isfinite(values)
        if int(valid.sum()) < 3:
            continue
        x, y = wavelengths[valid], values[valid]
        order = np.argsort(x, kind="stable")
        x, y = x[order], y[order]
        keep = np.concatenate(([True], np.diff(x) != 0))
        x, y = x[keep], y[keep]
        label = _unique_label(headers[value_index], used_labels)
        used_labels.add(label)
        series.append(SpectrumSeries(label, x, y, _value_unit(label)))
    if not series:
        raise ValueError("파장과 짝을 이루는 측정값 열을 찾지 못했습니다.")
    return SpectrumFile(path, encoding, delimiter, tuple(series))


def load_spectrum_file(filename: str | Path) -> SpectrumFile:
    path = Path(filename)
    if path.suffix.casefold() == ".xlsx":
        return load_spectrum_xlsx(path)
    return load_spectrum_csv(path)


def nearest_point(series: SpectrumSeries, wavelength_nm: float) -> tuple[float, float, int]:
    index = int(np.abs(series.wavelength_nm - wavelength_nm).argmin())
    return float(series.wavelength_nm[index]), float(series.values[index]), index


def summarize_region(
    series: SpectrumSeries, center_nm: float, half_width_nm: float
) -> RegionSummary:
    if half_width_nm <= 0:
        raise ValueError("영역 반폭은 0보다 커야 합니다.")
    mask = np.abs(series.wavelength_nm - center_nm) <= half_width_nm
    indices = np.flatnonzero(mask)
    if not indices.size:
        raise ValueError(f"{center_nm:g} nm 주변에 데이터가 없습니다.")
    local_values = series.values[indices]
    maximum_index = int(indices[int(np.argmax(local_values))])
    minimum_index = int(indices[int(np.argmin(local_values))])
    nearest_wavelength, nearest, _ = nearest_point(series, center_nm)
    return RegionSummary(
        center_nm=float(center_nm),
        half_width_nm=float(half_width_nm),
        count=int(indices.size),
        nearest_wavelength_nm=nearest_wavelength,
        nearest_value=nearest,
        maximum_wavelength_nm=float(series.wavelength_nm[maximum_index]),
        maximum_value=float(series.values[maximum_index]),
        minimum_wavelength_nm=float(series.wavelength_nm[minimum_index]),
        minimum_value=float(series.values[minimum_index]),
        mean_value=float(np.mean(local_values)),
    )


def export_series_csv(
    filename: str | Path,
    series: SpectrumSeries,
    minimum_nm: float | None = None,
    maximum_nm: float | None = None,
) -> int:
    lower = -np.inf if minimum_nm is None else float(minimum_nm)
    upper = np.inf if maximum_nm is None else float(maximum_nm)
    if lower > upper:
        lower, upper = upper, lower
    mask = (series.wavelength_nm >= lower) & (series.wavelength_nm <= upper)
    indices = np.flatnonzero(mask)
    if not indices.size:
        raise ValueError("선택 범위에 내보낼 데이터가 없습니다.")
    value_header = series.label
    if series.value_unit and series.value_unit.casefold() not in value_header.casefold():
        value_header = f"{value_header} ({series.value_unit})"
    with Path(filename).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["Wavelength (nm)", value_header])
        for index in indices:
            writer.writerow(
                [f"{series.wavelength_nm[index]:.6g}", f"{series.values[index]:.10g}"]
            )
    return int(indices.size)


def absorbance_from_transmittance(series: SpectrumSeries) -> SpectrumSeries:
    values = series.values.astype(float)
    percent_input = "%" in series.value_unit or "%t" in series.label.casefold()
    if not percent_input and float(np.nanmax(values)) > 1.5:
        percent_input = True
    fraction = values / 100.0 if percent_input else values
    if np.any(fraction <= 0):
        raise ValueError("0 이하의 투과율은 흡광도로 변환할 수 없습니다.")
    absorbance = -np.log10(fraction)
    return SpectrumSeries(
        label=f"{series.label} → 흡광도",
        wavelength_nm=series.wavelength_nm.copy(),
        values=absorbance,
        value_unit="A.U.",
    )


def is_transmittance(series: SpectrumSeries) -> bool:
    text = f"{series.label} {series.value_unit}".casefold()
    return any(word in text for word in ("transmit", "투과", "%t"))
